import os
import sys
import datetime
import difflib
import openpyxl

# Ensure repo root is on sys.path
repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if repo_root not in sys.path:
    sys.path.insert(0, repo_root)

import db

EXCEL_PATH = r"C:\Users\singarirangaprasad\Downloads\Surendra_Sales.xlsx"

CANONICAL_PRODUCT_MAP = {
    "king lights": "King Lights",
    "king lights": "King Lights",
    "king ligths": "King Lights",
    "king lights": "King Lights",
    "lights": "King Lights",
    "king red": "King Red",
    "king red": "King Red",
    "kiing red": "King Red",
    "kin red": "King Red",
    "king": "King Red",
    "ice brust": "Ice Burst",
    "ice burst": "Ice Burst",
    "u.milds": "Ultra Milds",
    "ultra milds": "Ultra Milds",
    "wilds": "Ultra Milds",
    "slk": "Sleeks",
    "slk delux": "Sleeks",
    "sleeks": "Sleeks",
    "filters": "Filter",
    "filter": "Filter",
    "connect": "Connect",
    "indie mint": "Indie Mint",
    "milds": "Milds",
    "define pan": "Define Paan",
    "define paan": "Define Paan",
    "advance": "Advance",
    "advance compact": "Advance Compact",
    "compact": "Compact",
    "double burst": "Double Burst",
    "navy cut": "Navy Cut",
    "wills": "Wills",
    "birstol": "Bristol",
    "classic clove": "Classic Clove",
    "clove": "Classic Clove",
    "american clove": "American Clove",
    "american l.i.t": "American L.I.T",
    "amrica l.i.t": "American L.I.T",
    "amrican l.i.t": "American L.I.T",
    "american n.y.cool": "American N.Y.Cool",
    "america n.y cool": "American N.Y.Cool",
    "n.y. cool": "American N.Y.Cool",
    "players red": "Players Red",
    "players blue": "Players Blue",
    "players": "Players",
    "social 2 pod": "Social 2 Pod",
    "socail-2 pod": "Social 2 Pod",
    "scoial 2 - pod": "Social 2 Pod",
    "twin pod": "Twin Pod",
    "twinpod": "Twin Pod",
    "mix pod": "Mix Pod",
    "mixpod": "Mix Pod",
    "shift define": "Shift Define",
    "shift": "Shift",
    "shitf": "Shift",
    "galaxy": "Galaxy",
    "gold": "Gold",
    "malabar gold": "Malabar Gold",
    "fine touch": "Fine Touch",
    "forest": "Forest",
    "fuse beyond": "Fuse Beyond",
    "club lit": "Club LIT",
    "neo smart": "Neo Smart",
    "scissors": "Scissors",
}

def classify_entry(item_str, sale_tot, purch_tot, cash, online, due):
    low = item_str.lower().strip()

    # Payments to owner
    if any(k in low for k in ["paid to prasad", "cash paid to prasad", "online paid to prasad", "paid to ranga", "paid cash to prasad", "paid online to prasad", "paid cash to ranga", "paid online to manasa", "paid to manasa", "paid to b reddy"]):
        return "PAYMENT_TO_OWNER"
    # Payments from owner
    if any(k in low for k in ["received from prasad", "cash received from prasad", "received from b reddy"]):
        return "PAYMENT_FROM_OWNER"
    # Credit Card Bills
    if any(k in low for k in ["credit card bill", "cc bill", "paid credit card"]):
        return "BILL_PAYMENT"
    # Expenses / Charges
    if any(k in low for k in ["auto charges", "train ticket", "yulu bike", "metro", "charges", "chares", "ticket", "trip"]) or " to " in low:
        return "EXPENSE"
    # Dues / Collections
    if any(k in low for k in ["received due", "received balance", "received cash"]):
        return "DUE_RECEIVED"
    # Adjustments
    if any(k in low for k in ["adjusted for the due", "excess paid", "paid excess", "wallet"]):
        return "ADJUSTMENT"
    # ATM transactions
    if "atm" in low:
        return "ADJUSTMENT"
    # Purchases
    if "purchase" in low or "purhcase" in low or (purch_tot > 0 and sale_tot == 0):
        return "PURCHASE"
    # Default to Sale
    return "SALE"

def normalize_product(item_str, catalogue_names):
    low = item_str.lower().strip()
    if low in CANONICAL_PRODUCT_MAP:
        return CANONICAL_PRODUCT_MAP[low]
    # Fuzzy match with catalogue names
    matches = difflib.get_close_matches(item_str, catalogue_names, n=1, cutoff=0.75)
    if matches:
        return matches[0]
    return item_str.strip()

def run_migration():
    print(f"Loading workbook: {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    if "Surendra" not in wb.sheetnames:
        print("Error: 'Surendra' sheet not found in workbook.")
        return

    ws = wb["Surendra"]
    print(f"Sheet 'Surendra' loaded. Rows: {ws.max_row}, Columns: {ws.max_column}")

    conn = db.get_connection()
    cur = conn.cursor()

    # Get existing catalogue product names for fuzzy matching
    cur.execute("SELECT name FROM products")
    catalogue_names = [r["name"] for r in cur.fetchall()]

    # Clear existing ledger entries if any
    cur.execute("DELETE FROM salesperson_ledger")
    conn.commit()

    rows_to_insert = []
    type_counts = {}
    total_excel_cash = 0.0
    total_excel_online = 0.0
    total_excel_due = 0.0

    for r in range(3, ws.max_row + 1):
        counterparty_raw = ws.cell(row=r, column=2).value # Col B
        date_raw = ws.cell(row=r, column=3).value         # Col C
        item_raw = ws.cell(row=r, column=4).value         # Col D
        qty_raw = ws.cell(row=r, column=5).value          # Col E
        rate_raw = ws.cell(row=r, column=9).value         # Col I (Sale Rate)
        purch_rate_raw = ws.cell(row=r, column=6).value   # Col F
        sale_tot_raw = ws.cell(row=r, column=10).value    # Col J
        purch_tot_raw = ws.cell(row=r, column=8).value    # Col H
        cash_raw = ws.cell(row=r, column=11).value        # Col K
        online_raw = ws.cell(row=r, column=12).value      # Col L
        due_raw = ws.cell(row=r, column=13).value         # Col M

        # Skip completely empty rows
        if not any([counterparty_raw, date_raw, item_raw, sale_tot_raw, purch_tot_raw, cash_raw, online_raw, due_raw]):
            continue

        # Format Date
        if isinstance(date_raw, (datetime.datetime, datetime.date)):
            entry_date = date_raw.strftime("%Y-%m-%d")
        elif isinstance(date_raw, str) and date_raw.strip():
            entry_date = date_raw.strip()[:10]
        else:
            entry_date = "2026-03-05" # Fallback if missing

        counterparty = str(counterparty_raw).strip() if counterparty_raw is not None else None
        item_desc = str(item_raw).strip() if item_raw is not None else "General Transaction"

        qty = float(qty_raw) if qty_raw is not None and isinstance(qty_raw, (int, float)) else 0.0
        sale_tot = float(sale_tot_raw) if sale_tot_raw is not None and isinstance(sale_tot_raw, (int, float)) else 0.0
        purch_tot = float(purch_tot_raw) if purch_tot_raw is not None and isinstance(purch_tot_raw, (int, float)) else 0.0
        rate = float(rate_raw) if rate_raw is not None and isinstance(rate_raw, (int, float)) else (float(purch_rate_raw) if purch_rate_raw is not None and isinstance(purch_rate_raw, (int, float)) else 0.0)

        cash = float(cash_raw) if cash_raw is not None and isinstance(cash_raw, (int, float)) else 0.0
        online = float(online_raw) if online_raw is not None and isinstance(online_raw, (int, float)) else 0.0
        due = float(due_raw) if due_raw is not None and isinstance(due_raw, (int, float)) else 0.0

        total_excel_cash += cash
        total_excel_online += online
        total_excel_due += due

        entry_type = classify_entry(item_desc, sale_tot, purch_tot, cash, online, due)
        type_counts[entry_type] = type_counts.get(entry_type, 0) + 1

        # Normalize product if SALE or PURCHASE
        if entry_type in ("SALE", "PURCHASE"):
            item_desc = normalize_product(item_desc, catalogue_names)

        total_amount = sale_tot if sale_tot > 0 else (purch_tot if purch_tot > 0 else abs(cash) + abs(online))

        rows_to_insert.append((
            "Surendra",
            entry_date,
            entry_type,
            counterparty,
            item_desc,
            round(qty, 2),
            round(rate, 2),
            round(total_amount, 2),
            round(cash, 2),
            round(online, 2),
            round(due, 2),
            f"Imported from Excel Row {r}"
        ))

    print(f"\nExtracted {len(rows_to_insert)} valid rows from Excel.")
    print("Classified entry type counts:")
    for t, c in sorted(type_counts.items(), key=lambda x: -x[1]):
        print(f"  {t:20s}: {c:5d}")

    # Insert in batches
    cur.executemany("""
        INSERT INTO salesperson_ledger (
            salesperson, entry_date, entry_type, counterparty, item_description,
            quantity, unit_rate, total_amount, cash_amount, online_amount, due_amount, notes
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, rows_to_insert)
    conn.commit()

    # Query back summary
    cur.execute("""
        SELECT 
            COUNT(*) as count,
            COALESCE(SUM(cash_amount), 0.0) as cash,
            COALESCE(SUM(online_amount), 0.0) as online,
            COALESCE(SUM(due_amount), 0.0) as due,
            COALESCE(SUM(total_amount), 0.0) as total
        FROM salesperson_ledger
    """)
    summary = cur.fetchone()
    conn.close()

    print("\n=== Verification Summary ===")
    print(f"Total Rows Inserted: {summary['count']}")
    print(f"Total Cash Balance : Rs. {summary['cash']:,.2f}  (Excel: Rs. 29,935.00)")
    print(f"Total Online Balance: Rs. {summary['online']:,.2f}  (Excel: -Rs. 94,227.32)")
    print(f"Total Due Balance   : Rs. {summary['due']:,.2f}  (Excel: Rs. 0.00)")
    print(f"Net Salesperson Position: Rs. {summary['cash'] + summary['online']:,.2f}  (Excel: -Rs. 64,292.32)")

if __name__ == "__main__":
    run_migration()
